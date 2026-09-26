# report.py
# -*- coding: utf-8 -*-
"""
用 python-docx 生成 Word 报告。

覆盖：描述统计 / 正态性检验 / t 检验 / 卡方 / ANOVA / 相关性 / 配对 t /
非参数检验（Mann-Whitney、Wilcoxon、Kruskal-Wallis、Friedman）/
事后多重比较 / 单样本 t / 比例检验 / McNemar /
逻辑回归 / 多元线性回归 / 信度分析 Cronbach's α / 图形。
"""

import io

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

from stats_utils import format_p


def _set_font(run, name="宋体", size=None):
    """设置中文字体。"""
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)


def _apa_borders(table, header_rows=1):
    """给表格套 APA 三线表：顶线 + 表头下栏目线 + 底线，无竖线、无内部横线。"""
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    spec = [("top", 12), ("bottom", 12)]                       # 粗顶线 / 粗底线
    for edge, sz in spec:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    for edge in ("left", "right", "insideH", "insideV"):     # 无竖线 / 无内部横线
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        borders.append(el)
    tblPr.append(borders)
    # 表头行底边：栏目（header）线
    for cell in table.rows[header_rows - 1].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcB = tcPr.find(qn("w:tcBorders"))
        if tcB is None:
            tcB = OxmlElement("w:tcBorders")
            tcPr.append(tcB)
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")
        tcB.append(bottom)


def _add_dataframe_table(doc, df, title=None):
    """把 DataFrame 以 APA 三线表插入 Word 报告。"""
    if title:
        doc.add_heading(title, level=3)

    df_show = df.reset_index()
    table = doc.add_table(rows=1, cols=len(df_show.columns))
    _apa_borders(table)

    hdr = table.rows[0].cells
    for i, col in enumerate(df_show.columns):
        hdr[i].text = str(col)
        for p in hdr[i].paragraphs:
            for r in p.runs:
                _set_font(r, "宋体", 9)

    for _, row in df_show.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if pd.isna(val):
                text = "—"
            elif isinstance(val, float):
                text = f"{val:.4f}"
            else:
                text = str(val)
            cells[i].text = text
            for p in cells[i].paragraphs:
                for r in p.runs:
                    _set_font(r, "宋体", 9)


def _add_figure(doc, fig, caption):
    """把 matplotlib 图插入 Word。"""
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    buf.seek(0)
    doc.add_picture(buf, width=Inches(5.8))
    doc.add_paragraph(caption).alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf.close()


def _add_sec(doc, counter, title):
    """添加带序号的一级章节。"""
    counter[0] += 1
    doc.add_heading(f"{counter[0]}、{title}", level=1)


def _fmt_num(x, fmt=".4f"):
    if x is None or (isinstance(x, float) and (np.isnan(x) or pd.isna(x))):
        return "—"
    return f"{x:{fmt}}"


def build_word_report(df, data_name, desc_table=None, norm_table=None,
                      ttest=None, reg=None, figures=None, alpha=0.05,
                      chi2=None, anova=None, corr=None, paired=None,
                      mw=None, wsr=None, kw=None, fried=None, post=None,
                      osamp=None, prop=None, mcn=None, logreg=None,
                      multireg=None, cron=None):
    """生成 Word 报告，返回 bytes。"""
    doc = Document()
    sec = [0]

    # 标题
    title = doc.add_heading("统计分析报告", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"数据文件：{data_name}")
    doc.add_paragraph(f"生成时间：{pd.Timestamp.now():%Y-%m-%d %H:%M:%S}")
    doc.add_paragraph(f"显著性水平 α = {alpha}")

    # 一、数据概况
    _add_sec(doc, sec, "数据概况")
    doc.add_paragraph(
        f"观测数：{df.shape[0]} 行；变量数：{df.shape[1]} 列；"
        f"缺失值总数：{int(df.isna().sum().sum())}。"
    )
    info = pd.DataFrame({
        "数据类型": df.dtypes.astype(str),
        "非缺失数": df.notna().sum(),
        "缺失数": df.isna().sum(),
        "缺失比例(%)": (df.isna().mean() * 100).round(2),
        "唯一值个数": df.nunique(dropna=True),
    })
    _add_dataframe_table(doc, info, "变量信息")

    # 二、描述统计
    if desc_table is not None and not desc_table.empty:
        _add_sec(doc, sec, "描述统计")
        _add_dataframe_table(doc, desc_table, "描述统计表")

    # 三、正态性检验
    if norm_table is not None and not norm_table.empty:
        _add_sec(doc, sec, "正态性检验")
        norm_show = norm_table.copy()
        if "p 值" in norm_show.columns:
            norm_show["p 值"] = norm_show["p 值"].apply(
                lambda p: "<0.001" if pd.notna(p) and p < 0.001
                else (f"{p:.4f}" if pd.notna(p) else "—")
            )
        _add_dataframe_table(doc, norm_show, "正态性检验结果")
        doc.add_paragraph("注：p > α 时不拒绝正态假设（可视为近似正态）。")

    # 四、t 检验
    if ttest:
        _add_sec(doc, sec, "独立样本 t 检验")
        r = ttest
        doc.add_paragraph(f"因变量：{r['因变量']}；分组变量：{r['分组变量']}。")
        doc.add_paragraph(
            f"{r['水平1']}：n={r['n1']}，均值={r['mean1']:.4f}，标准差={r['std1']:.4f}。"
        )
        doc.add_paragraph(
            f"{r['水平2']}：n={r['n2']}，均值={r['mean2']:.4f}，标准差={r['std2']:.4f}。"
        )
        doc.add_paragraph(
            f"Levene 方差齐性检验：统计量={r['levene_stat']:.4f}，"
            f"p={r['levene_p']:.4f}，方差{'齐性' if r['equal_var'] else '不齐性'}。"
        )
        doc.add_paragraph(
            f"检验方法：{r['检验方法']}；t={r['t']:.4f}，"
            f"df={r['df']:.2f}，p={r['p']:.4f}。"
        )
        doc.add_paragraph(
            f"均值差={r['mean_diff']:.4f}，95% CI=[{r['ci_low']:.4f}, {r['ci_high']:.4f}]。"
        )
        doc.add_paragraph(
            f"Cohen's d={r['cohens_d']:.4f}（{r['effect_label']}）。"
        )
        doc.add_paragraph(
            f"Mann-Whitney U（非参备选）：U={r['mannwhitney_u']:.1f}，"
            f"p={r['mannwhitney_p']:.4f}。"
        )
        doc.add_paragraph(f"结论：{r['结论']}")

    # 五、卡方检验
    if chi2:
        _add_sec(doc, sec, "卡方独立性检验")
        r = chi2
        doc.add_paragraph(f"检验变量：{r['列1']} 与 {r['列2']}；样本量={r['样本量']}。")
        cv = r["Cramér's V"]
        doc.add_paragraph(
            f"χ²={r['卡方统计量']:.4f}，df={r['自由度']}，"
            f"p={format_p(r['p 值'])}，Cramér's V={cv:.4f}。"
        )
        _add_dataframe_table(doc, r["列联表"], "列联表（观测频数）")
        doc.add_paragraph(f"结论：{r['结论']}")

    # 六、方差分析 ANOVA
    if anova:
        _add_sec(doc, sec, "单因素方差分析 ANOVA")
        r = anova
        doc.add_paragraph(f"因变量：{r['因变量']}；分组变量：{r['分组变量']}（{r['组数']} 组）。")
        gm = "；".join(f"{k}={v:.4f}" for k, v in r["组均值"].items())
        doc.add_paragraph(f"各组均值：{gm}。")
        dfree = r["自由度"]
        doc.add_paragraph(
            f"F={r['F 统计量']:.4f}，df=({dfree[0]}, {dfree[1]})，p={format_p(r['p 值'])}。"
        )
        f_eff = r.get("Cohen's f")
        doc.add_paragraph(
            f"效应量：η²={_fmt_num(r.get('eta²'))}，ω²={_fmt_num(r.get('omega²'))}，"
            f"Cohen's f={_fmt_num(f_eff)}。"
        )
        doc.add_paragraph(f"结论：{r['结论']}")

    # 七、相关性检验
    if corr:
        _add_sec(doc, sec, "相关性检验")
        r = corr
        doc.add_paragraph(f"方法：{r['方法']}；变量：{r['X']} 与 {r['Y']}；样本量={r['样本量']}。")
        doc.add_paragraph(
            f"相关系数 r={r['相关系数 r']:.4f}，p={format_p(r['p 值'])}，"
            f"相关强度：{r['强度']}。"
        )
        doc.add_paragraph(f"结论：{r['结论']}")

    # 八、配对 t 检验
    if paired:
        _add_sec(doc, sec, "配对 t 检验")
        r = paired
        doc.add_paragraph(f"前测列：{r['前测列']}；后测列：{r['后测列']}；样本对={r['样本对']}。")
        cohens_d = r["Cohen's d"]
        doc.add_paragraph(
            f"均值差(后-前)={r['均值差(后-前)']:.4f}，t={r['t 统计量']:.4f}，"
            f"df={r['自由度']}，p={format_p(r['p 值'])}，Cohen's d={cohens_d:.4f}。"
        )
        if not pd.isna(r.get("Wilcoxon p", np.nan)):
            doc.add_paragraph(f"Wilcoxon 非参备选：p={r['Wilcoxon p']:.4f}。")
        doc.add_paragraph(f"结论：{r['结论']}")

    # 九、非参数检验（M-W / Wilcoxon / Kruskal / Friedman）
    nonparam = []
    if mw:
        r = mw
        nonparam.append((
            "Mann-Whitney U 检验",
            f"因变量：{r['因变量']}；分组：{r['分组变量']}。"
            f"{r['水平1']} 中位数={r['中位数1']:.4f}（n={r['n1']}），"
            f"{r['水平2']} 中位数={r['中位数2']:.4f}（n={r['n2']}）。"
            f"U={r['U 统计量']:.1f}，p={format_p(r['p 值'])}，"
            f"秩双列相关={_fmt_num(r['效应量(秩双列相关)'])}。结论：{r['结论']}",
        ))
    if wsr:
        r = wsr
        nonparam.append((
            "Wilcoxon 符号秩检验",
            f"前测：{r['前测列']}；后测：{r['后测列']}；样本对={r['样本对']}。"
            f"中位数差(后-前)={r['中位数差(后-前)']:.4f}，W={r['W 统计量']:.1f}，"
            f"p={format_p(r['p 值'])}，效应量 r={_fmt_num(r['效应量 r'])}。结论：{r['结论']}",
        ))
    if kw:
        r = kw
        meds = "；".join(f"{k}={v:.4f}" for k, v in r["组中位数"].items())
        nonparam.append((
            "Kruskal-Wallis H 检验",
            f"因变量：{r['因变量']}；分组：{r['分组变量']}（{r['组数']} 组，样本量={r['样本量']}）。"
            f"各组中位数：{meds}。H={r['H 统计量']:.4f}，df={r['自由度']}，"
            f"p={format_p(r['p 值'])}，ε²={_fmt_num(r['效应量 ε²'])}。结论：{r['结论']}",
        ))
    if fried:
        r = fried
        meds = "；".join(f"{k}={v:.4f}" for k, v in r["各条件中位数"].items())
        kendall_w = r["Kendall's W"]
        nonparam.append((
            "Friedman 检验",
            f"测量条件：{'、'.join(r['测量条件'])}；样本量={r['样本量']}。"
            f"各条件中位数：{meds}。Q={r['Q 统计量']:.4f}，df={r['自由度']}，"
            f"p={format_p(r['p 值'])}，Kendall's W={_fmt_num(kendall_w)}。结论：{r['结论']}",
        ))
    if nonparam:
        _add_sec(doc, sec, "非参数检验")
        for name, text in nonparam:
            doc.add_heading(name, level=3)
            doc.add_paragraph(text)

    # 十、事后多重比较
    if post:
        _add_sec(doc, sec, "事后多重比较")
        r = post
        doc.add_paragraph(
            f"方法：{r['方法']}；因变量：{r['因变量']}；分组变量：{r['分组变量']}；"
            f"共比较 {r['比较对数']} 对。"
        )
        _add_dataframe_table(doc, r["比较结果表"], f"成对比较结果（{r['方法']}）")
        doc.add_paragraph(f"结论：{r['结论']}")

    # 十一、单样本 t / 比例 / McNemar
    if osamp:
        _add_sec(doc, sec, "单样本 t 检验")
        r = osamp
        ci = r["95%CI"]
        cohens_d = r["Cohen's d"]
        doc.add_paragraph(
            f"变量：{r['变量']}；样本量={r['样本量']}；检验值 μ0={r['检验值 μ0']}。"
            f"均值={r['均值']:.4f}，t={r['t 统计量']:.4f}，df={r['自由度']}，"
            f"p={format_p(r['p 值'])}，95%CI=[{ci[0]:.4f}, {ci[1]:.4f}]，"
            f"Cohen's d={_fmt_num(cohens_d)}（{r['效应量评价']}）。结论：{r['结论']}"
        )
    if prop:
        _add_sec(doc, sec, "二项比例检验")
        r = prop
        ci = r["95%CI(Wilson)"]
        doc.add_paragraph(
            f"变量：{r['变量']}；成功水平={r['成功水平']}；样本量={r['样本量']}。"
            f"样本比例={r['样本比例']:.4f}（成功 {r['成功数']} 例），检验比例 p0={r['检验比例 p0']}。"
            f"精确 p（Binomial）={format_p(r['精确 p 值(Binomial)'])}，"
            f"95%CI(Wilson)=[{ci[0]:.4f}, {ci[1]:.4f}]。结论：{r['结论']}"
        )
    if mcn:
        _add_sec(doc, sec, "McNemar 配对卡方检验")
        r = mcn
        doc.add_paragraph(
            f"变量1：{r['变量1']}；变量2：{r['变量2']}；样本量={r['样本量']}。"
            f"McNemar χ²(校正)={r['McNemar χ²(校正)']:.4f}，p={format_p(r['p 值'])}，"
            f"精确 p（Binomial）={format_p(r['精确 p 值(Binomial)'])}，"
            f"不一致对数={r['不一致对数']}，一致率={r['一致率']:.4f}。"
        )
        _add_dataframe_table(doc, r["列联表"], "配对列联表")
        doc.add_paragraph(f"结论：{r['结论']}")

    # 十二、回归分析（简单 / 逻辑 / 多元）
    if reg:
        _add_sec(doc, sec, "简单线性回归")
        r = reg
        doc.add_paragraph(f"回归方程：{r['equation']}，n={r['n']}。")
        doc.add_paragraph(
            f"R²={r['r2']:.4f}，调整 R²={r['adj_r2']:.4f}，"
            f"F={r['f_stat']:.4f}，p={r['f_p']:.4f}。"
        )
        doc.add_paragraph(
            f"常数项 β₀={r['b0']:.4f}（SE={r['se0']:.4f}, "
            f"t={r['t0']:.4f}, p={r['p0']:.4f}）。"
        )
        doc.add_paragraph(
            f"斜率 β₁={r['b1']:.4f}（SE={r['se1']:.4f}, "
            f"t={r['t1']:.4f}, p={r['p1']:.4f}）。"
        )
        doc.add_paragraph(
            f"Pearson r={r['pearson_r']:.4f}，p={r['pearson_p']:.4f}。"
        )
        doc.add_paragraph(f"结论：{r['结论']}")
    if logreg:
        _add_sec(doc, sec, "二分类逻辑回归")
        r = logreg
        doc.add_paragraph(
            f"因变量：{r['因变量']}（成功水平={r['成功水平']}）；"
            f"自变量：{'、'.join(r['自变量'])}；样本量={r['样本量']}。"
        )
        doc.add_paragraph(
            f"McFadden R²={_fmt_num(r['McFadden R²'])}，整体似然比 p={format_p(r['整体似然比 p'])}，"
            f"准确率={_fmt_num(r['准确率'])}，AUC={_fmt_num(r['AUC'])}。"
        )
        _add_dataframe_table(doc, r["系数表"], "逻辑回归系数表")
        doc.add_paragraph(f"结论：{r['结论']}")
    if multireg:
        _add_sec(doc, sec, "多元线性回归")
        r = multireg
        doc.add_paragraph(
            f"因变量：{r['因变量']}；自变量：{'、'.join(r['自变量'])}；样本量={r['样本量']}。"
        )
        doc.add_paragraph(
            f"R²={r['R²']:.4f}，调整 R²={r['调整 R²']:.4f}，F={r['F 统计量']:.4f}，"
            f"p={format_p(r['整体 p 值'])}，Durbin-Watson={r['Durbin-Watson']:.4f}，"
            f"RMSE={_fmt_num(r['RMSE'])}，AIC={_fmt_num(r['AIC'], '.1f')}，BIC={_fmt_num(r['BIC'], '.1f')}。"
        )
        _add_dataframe_table(doc, r["系数表"], "多元回归系数表")
        _add_dataframe_table(doc, r["VIF 表"], "共线性诊断（VIF）")
        doc.add_paragraph(f"结论：{r['结论']}")

    # 十三、信度分析
    if cron:
        _add_sec(doc, sec, "信度分析（Cronbach's α）")
        r = cron
        alpha_val = r["Cronbach's α"]
        doc.add_paragraph(f"样本量={r['样本量']}，题项数={r['题项数']}。")
        doc.add_paragraph(
            f"Cronbach's α={_fmt_num(alpha_val)}，标准化 α={_fmt_num(r['标准化 α'])}，"
            f"平均项间相关={_fmt_num(r['平均项间相关'])}。"
        )
        _add_dataframe_table(doc, r["删除题项后 α"], "删除各题项后的 α")
        doc.add_paragraph(f"结论：{r['结论']}")

    # 图形
    if figures:
        _add_sec(doc, sec, "图形展示")
        for fig, caption in figures:
            _add_figure(doc, fig, caption)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.getvalue()